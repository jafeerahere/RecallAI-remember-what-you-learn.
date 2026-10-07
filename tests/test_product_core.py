from io import BytesIO
from datetime import datetime, timezone
import pytest
from core.ingestion import extract_document, from_text, source_location
from core.mastery import mastery_score, schedule_review, mastery_label
from database.repositories import Repository

def test_formats_and_locations():
    sentence = 'Attention assigns different weights to tokens in a sequence.'
    txt = extract_document(sentence.encode(), 'Attention.txt')
    assert txt.title == 'Attention' and txt.file_type == 'TXT'
    from docx import Document
    doc = Document(); doc.add_heading('Attention', 1); doc.add_paragraph(sentence)
    stream = BytesIO(); doc.save(stream)
    material = extract_document(stream.getvalue(), 'lesson.docx')
    assert source_location(sentence, material.segments)['section'] == 'Attention'
    from pptx import Presentation
    slides = Presentation(); slide = slides.slides.add_slide(slides.slide_layouts[1])
    slide.shapes.title.text = 'Attention'; slide.placeholders[1].text = sentence
    stream = BytesIO(); slides.save(stream)
    assert extract_document(stream.getvalue(),'lesson.pptx').segments[0]['slide'] == 1
    from reportlab.pdfgen import canvas
    stream = BytesIO(); pdf = canvas.Canvas(stream); pdf.drawString(40,750,sentence); pdf.save()
    material = extract_document(stream.getvalue(),'lesson.pdf')
    assert source_location(sentence, material.segments)['page'] == 1

def test_invalid_extraction():
    for payload, filename in [(b'','empty.txt'),(b'broken','bad.pdf'),(b'broken','bad.docx'),(b'x','bad.exe'),(b'\xff','bad.txt')]:
        with pytest.raises(ValueError):
            extract_document(payload, filename)
    from pypdf import PdfWriter
    writer = PdfWriter(); writer.add_blank_page(300,300)
    stream = BytesIO(); writer.write(stream)
    with pytest.raises(ValueError,match='No readable text'):
        extract_document(stream.getvalue(),'scan.pdf')
    writer.encrypt('password'); stream = BytesIO(); writer.write(stream)
    with pytest.raises(ValueError,match='password-protected'):
        extract_document(stream.getvalue(),'locked.pdf')

def test_mastery_and_schedule():
    assert mastery_score([]) == 0
    assert mastery_score([{'value':1}]) == 20
    assert mastery_score([{'value':1}]*5) == 100
    assert mastery_score([{'value':1}]*5 + [{'value':0}]*3) < 40
    assert mastery_label(90) == 'Mastered'
    now = datetime(2026,10,7,tzinfo=timezone.utc)
    again = schedule_review({},'Again',now)
    easy = schedule_review({},'Easy',now)
    assert again['next_review'] == '2026-10-07T00:10:00+00:00'
    assert easy['next_review'] == '2026-10-11T00:00:00+00:00'
    with pytest.raises(ValueError):
        schedule_review({},'Invalid')

def test_greeting_uses_time():
    from services.analytics import greeting
    assert greeting(datetime(2026,10,7,8,tzinfo=timezone.utc)) == 'Good morning'
    assert greeting(datetime(2026,10,7,14,tzinfo=timezone.utc)) == 'Good afternoon'
    assert greeting(datetime(2026,10,7,20,tzinfo=timezone.utc)) == 'Good evening'

def test_persistence_search_and_restore(tmp_path):
    path = tmp_path / 'library.sqlite3'
    repo = Repository(path)
    material = from_text('Transformers use attention to represent relationships between tokens.','Transformers')
    identifier, duplicate = repo.save_document(material)
    assert not duplicate
    assert repo.save_document(material) == (identifier, True)
    fresh = Repository(path)
    assert fresh.document(identifier)['text'] == material.text
    assert fresh.search('attention')[0]['document_id'] == identifier
    assert not fresh.search("%' OR 1=1 --")
    fresh.rename(identifier,'My transformers notes')
    assert repo.document(identifier)['title'] == 'My transformers notes'
    fresh.save_preferences({'depth':'Deep'})
    assert repo.preferences()['depth'] == 'Deep'
    fresh.trash(identifier)
    assert repo.document(identifier) is None and not repo.documents()
    fresh.restore(identifier)
    assert repo.document(identifier)
