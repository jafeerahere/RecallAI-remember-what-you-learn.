"""Extract learning material while retaining real source locations."""
from dataclasses import dataclass, asdict
from io import BytesIO
from pathlib import Path
import hashlib
import re
import zipfile

MAX_CHARACTERS = 250_000
MAX_UPLOAD_BYTES = 20 * 1024 * 1024

def normalize(text):
    return re.sub(r'\s+', ' ', text.replace('\x00', '')).strip()

@dataclass
class Material:
    title: str
    file_type: str
    text: str
    segments: list
    warnings: list

    @property
    def content_hash(self):
        import json
        payload = {'text': self.text, 'file_type': self.file_type, 'segments': self.segments}
        return hashlib.sha256(json.dumps(payload, sort_keys=True).encode()).hexdigest()

    def to_dict(self):
        return asdict(self)

def _finish(title, kind, segments, warnings=None):
    clean = [{**s, 'text': normalize(s['text'])} for s in segments if normalize(s['text'])]
    text = ' '.join(s['text'] for s in clean)
    if not text:
        raise ValueError('No readable text was found. Image-only documents need OCR before uploading.')
    if len(text) > MAX_CHARACTERS:
        raise ValueError('This material is too long. Split it into chapters under 250,000 characters.')
    return Material(title.strip()[:200] or 'Untitled material', kind, text, clean, warnings or [])

def from_text(text, title='Untitled material', file_type='Paste'):
    segments, section = [], 'Notes'
    for index, block in enumerate(re.split(r'\n\s*\n', text), 1):
        block = block.strip()
        if not block:
            continue
        lines = block.splitlines()
        if lines[0].startswith('#'):
            section = lines[0].lstrip('# ').strip()
            block = '\n'.join(lines[1:]) or section
        segments.append({'text': block, 'section': section, 'paragraph': index})
    return _finish(title, file_type, segments)

def _check_zip(payload):
    try:
        with zipfile.ZipFile(BytesIO(payload)) as archive:
            entries = archive.infolist()
            if len(entries) > 5000 or sum(i.file_size for i in entries) > 80 * 1024 * 1024:
                raise ValueError('This file expands to too much content. Upload a smaller document.')
    except zipfile.BadZipFile as exc:
        raise ValueError('This document is damaged or is not a valid Office file.') from exc

def extract_document(payload, filename):
    if not payload:
        raise ValueError('This file is empty. Choose a file that contains text.')
    if len(payload) > MAX_UPLOAD_BYTES:
        raise ValueError('Upload a file smaller than 20 MB.')
    extension = Path(filename).suffix.lower()
    title = Path(filename).stem
    if extension == '.txt':
        try:
            return from_text(payload.decode('utf-8-sig'), title, 'TXT')
        except UnicodeDecodeError as exc:
            raise ValueError('Save the TXT file with UTF-8 encoding, then upload it again.') from exc
    if extension not in ('.pdf', '.docx', '.pptx'):
        raise ValueError('Supported formats are PDF, DOCX, PPTX and TXT.')
    try:
        if extension == '.pdf':
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(payload))
            if reader.is_encrypted:
                raise ValueError('This PDF is password-protected. Upload an unlocked copy.')
            if len(reader.pages) > 500:
                raise ValueError('Split this PDF into files of at most 500 pages.')
            segments, missing = [], []
            for i, page in enumerate(reader.pages, 1):
                text = page.extract_text() or ''
                if not normalize(text):
                    missing.append(i)
                else:
                    segments.append({'text': text, 'page': i, 'section': f'Page {i}'})
            warnings = [f'No text on pages {", ".join(map(str, missing))}. Images on those pages were not read.'] if missing else []
            return _finish(title, 'PDF', segments, warnings)
        _check_zip(payload)
        if extension == '.docx':
            from docx import Document
            from docx.text.paragraph import Paragraph
            doc = Document(BytesIO(payload))
            section, segments, paragraph = 'Document', [], 0
            for block in doc.iter_inner_content():
                paragraph += 1
                if isinstance(block, Paragraph):
                    text = block.text
                    if block.style and block.style.name.startswith(('Heading', 'Title')):
                        section = text.strip() or section
                else:
                    text = '\n'.join(' | '.join(cell.text for cell in row.cells) for row in block.rows)
                segments.append({'text': text, 'section': section, 'paragraph': paragraph})
            return _finish(title, 'DOCX', segments)
        from pptx import Presentation
        presentation = Presentation(BytesIO(payload))
        segments = []
        def shape_text(shape):
            if hasattr(shape, 'shapes'):
                return '\n'.join(shape_text(child) for child in shape.shapes)
            if shape.has_text_frame:
                return shape.text
            if shape.has_table:
                return '\n'.join(' | '.join(cell.text for cell in row.cells) for row in shape.table.rows)
            return ''
        for i, slide in enumerate(presentation.slides, 1):
            heading = slide.shapes.title.text if slide.shapes.title else f'Slide {i}'
            text = '\n'.join(shape_text(shape) for shape in slide.shapes)
            segments.append({'text': text, 'slide': i, 'section': heading})
        return _finish(title, 'PPTX', segments)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError('Could not read this file. Check that it opens normally and upload a fresh copy.') from exc

def source_location(source, segments):
    """Never invent a page number; use extraction metadata only."""
    source = normalize(source)
    for segment in segments:
        if source in segment['text']:
            return {k: v for k, v in segment.items() if k != 'text'}
    return {'section': 'Document'}
